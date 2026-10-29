"""Verify actual PowerPoint-rendered PDF against content, plus slide geometry."""

import json
import re
from pathlib import Path

import pymupdf
from pptx import Presentation


def compact(text):
    return re.sub(r"[^a-z0-9]", "", text.lower())


def main(pdf_path="reports/release/presentation_powerpoint.pdf"):
    content = json.loads(Path("reports/release/presentation_content.json").read_text())
    prs = Presentation("docs/presentation_slides.pptx")
    pdf = pymupdf.open(pdf_path)
    assert len(prs.slides) == len(pdf) == len(content["slides"]) == 7
    rows = []
    for index, (slide, page, data) in enumerate(zip(prs.slides, pdf, content["slides"]), 1):
        expected = [data["title"], data["subtitle"]] + data["bullets"]
        expected += [value for row in data.get("table", []) for value in row]
        expected += data.get("flow", [])
        if "feedback" in data:
            expected.append(data["feedback"])
        actual = compact(page.get_text())
        missing = [text for text in expected if compact(text) not in actual]
        assert not missing, (index, missing)
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 10
            assert shape.top + shape.height <= prs.slide_height + 10
        spans = [s for block in page.get_text("dict")["blocks"] if "lines" in block
                 for line in block["lines"] for s in line["spans"]]
        footer = min(s["bbox"][1] for s in spans if "Synthetic MU economics" in s["text"])
        body = [s for s in spans if s["bbox"][1] > 130 and "Synthetic MU economics" not in s["text"]]
        assert max(s["bbox"][3] for s in body) < footer - 2, (index, "body/footer overlap")
        page.get_pixmap(matrix=pymupdf.Matrix(1.4, 1.4)).save(f"reports/release/browser/powerpoint_slide_{index}.png")
        rows.append(dict(slide=index, text_complete=True, shapes_within_canvas=True, footer_clear=True))
    result = dict(status="PASS", renderer="Local Microsoft PowerPoint PDF export", slides=rows,
                  talk_track_seconds=content["total_seconds"], candidate_rehearsal="Prepared; not claimed performed")
    Path("reports/release/presentation_verification.json").write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
